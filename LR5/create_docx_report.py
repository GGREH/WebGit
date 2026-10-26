import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_report():
    doc = docx.Document()
    
    # Встановлення параметрів сторінки (А4, поля: ліве 2.5 см, праве 1.5 см, верхнє 2.0 см, нижнє 2.0 см)
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.left_margin = Inches(0.98)   # ~2.5 cm
        section.right_margin = Inches(0.59)  # ~1.5 cm
        section.top_margin = Inches(0.79)    # ~2.0 cm
        section.bottom_margin = Inches(0.79) # ~2.0 cm

    # Базовий стиль для нормального тексту
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(14)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.25
    normal_style.paragraph_format.space_after = Pt(4)

    # ----------------------------------------------------
    # ТИТУЛЬНА СТОРІНКА
    # ----------------------------------------------------
    def add_p(text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, size=14, space_after=4, space_before=0, italic=False):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.line_spacing = 1.2
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        return p

    # Шапка
    add_p("КИЇВСЬКИЙ НАЦІОНАЛЬНИЙ УНІВЕРСИТЕТ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14)
    add_p("ІМЕНІ ТАРАСА ШЕВЧЕНКА", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14)
    add_p("ФАКУЛЬТЕТ ІНФОРМАЦІЙНИХ ТЕХНОЛОГІЙ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, space_before=10)
    add_p("КАФЕДРА ІНТЕЛЕКТУАЛЬНИХ ТЕХНОЛОГІЙ", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, space_before=10, space_after=60)

    # Заголовок роботи
    add_p("Лабораторна робота №5", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=18, space_after=10)
    add_p("з дисципліни «Проектування та розробка веб-застосувань»", align=WD_ALIGN_PARAGRAPH.CENTER, size=14, space_after=6)
    add_p("Тема Роботи: «Основи мови JavaScript: робота з об'єктом Date, масивами, циклами, об'єктами користувача, моделлю DOM та подіями»", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, space_after=14)
    add_p("Варіант №16 (Варіант 6)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=15, space_after=70)

    # Виконавець та викладач
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_meta.paragraph_format.space_after = Pt(70)
    p_meta.paragraph_format.line_spacing = 1.25
    r_meta = p_meta.add_run(
        "Виконав студент\n"
        "групи АнД - 31:\n"
        "Мамон Р.Д.\n\n"
        "Перевірила:\n"
        "Федусенко О.В."
    )
    r_meta.font.name = 'Times New Roman'
    r_meta.font.size = Pt(14)

    # Місто та рік
    add_p("Київ - 2026", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14)

    doc.add_page_break()

    # ----------------------------------------------------
    # ТЕМА, МЕТА, ЗАВДАННЯ
    # ----------------------------------------------------
    add_p("Тема", bold=True, size=16, space_before=10, space_after=6)
    add_p("Основи мови JavaScript: робота з об'єктом Date, масивами, циклами, об'єктами користувача, моделлю DOM та подіями.", space_after=12)

    add_p("Мета", bold=True, size=16, space_before=10, space_after=6)
    add_p("Розглянути можливості роботи з датою в JavaScript (об'єкт Date()), використання масивів та циклів у JavaScript, перевірки значень у коді скрипта. Навчитися створювати і використовувати об'єкти користувача в JavaScript з використанням моделі DOM та працювати з подіями.", space_after=12)

    add_p("Завдання", bold=True, size=16, space_before=10, space_after=6)
    add_p("Необхідно розробити та протестувати інтерактивний веб-застосунок на основі мови JavaScript за закріпленим варіантом (номер у групі 16, варіант індивідуального завдання 6):")

    add_p("1. Робота з масивами об'єктів та об'єктом Date (Завдання 1, Варіант 6):", bold=True, size=14, space_before=6)
    bullets_1 = [
        "Створити масив об'єктів про кінофільми з обов'язковими полями: назва фільму, кінотеатр, зал (допускається порожнє значення), дата початку прокату, сеанс.",
        "Реалізувати розрахунок тривалості прокату (кожен фільм демонструється у кінотеатрі рівно 10 днів від дати прем'єри).",
        "Реалізувати обчислення та наочне виведення інформації про те, скільки днів залишилося до показу фільмів, які фільми демонструються сьогодні, та скільки днів залишилося до завершення їх показу.",
        "Якщо зал не вказаний, вивести додаткове текстове повідомлення: «У кінотеатрі лише один зал».",
        "Здійснити кольорове маркування статусів прокату (активний показ, завершення прокату, показ закінчено, очікується прем'єра)."
    ]
    for b in bullets_1:
        p_b = doc.add_paragraph(style='List Bullet')
        p_b.paragraph_format.line_spacing = 1.25
        p_b.paragraph_format.space_after = Pt(3)
        r = p_b.add_run(b)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13.5)

    add_p("2. Маніпуляція вузлами DOM-дерева (Завдання 2.1, Варіант 6):", bold=True, size=14, space_before=8)
    bullets_2 = [
        "Розмістити на веб-сторінці блоки <div>, які містять текстову інформацію.",
        "Обійти всі вказані блоки за допомогою засобів об'єктної моделі документа (DOM) та динамічно сформувати впорядкований нумерований список (<ol>), елементами якого є текстовий вміст блоків."
    ]
    for b in bullets_2:
        p_b = doc.add_paragraph(style='List Bullet')
        p_b.paragraph_format.line_spacing = 1.25
        p_b.paragraph_format.space_after = Pt(3)
        r = p_b.add_run(b)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13.5)

    add_p("3. Обробка числової матриці та динамічна зміна стилів (Завдання 2.2, Варіант 6):", bold=True, size=14, space_before=8)
    bullets_3 = [
        "Задати на сторінці таблицю з числами (розмірністю 4х4), яка містить як додатні, так і від'ємні значення.",
        "При натисканні на кнопку зчитати значення комірок у двовимірний масив і виконати пошук екстремумів: окремо визначити максимальне та мінімальне додатне число, а також максимальне (найближче до 0) та мінімальне від'ємне число.",
        "Виокремити кольорами всі додатні та від'ємні числа в таблиці, а знайдені екстремальні значення виділити контрастними акцентними кольорами.",
        "(Додатково) Реалізувати обчислення та виведення сум чисел за кожним окремим стовпцем числової матриці."
    ]
    for b in bullets_3:
        p_b = doc.add_paragraph(style='List Bullet')
        p_b.paragraph_format.line_spacing = 1.25
        p_b.paragraph_format.space_after = Pt(3)
        r = p_b.add_run(b)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13.5)

    doc.add_page_break()

    # ----------------------------------------------------
    # ХІД РОБОТИ
    # ----------------------------------------------------
    add_p("Хід роботи", bold=True, size=16, space_before=10, space_after=8)

    p_links = doc.add_paragraph()
    p_links.paragraph_format.space_after = Pt(10)
    r1 = p_links.add_run("Посилання на GitHub репозиторій: ")
    r1.bold = True
    r2 = p_links.add_run("https://github.com/GGREH/WebGit.git\n")
    r2.font.color.rgb = RGBColor(0, 51, 153)
    r3 = p_links.add_run("Шлях до каталогу проекту: ")
    r3.bold = True
    r4 = p_links.add_run("https://github.com/GGREH/WebGit/tree/main/LR5")
    r4.font.color.rgb = RGBColor(0, 51, 153)

    add_p("Веб-застосунок побудовано за модульною архітектурою, яка включає такі складові:")
    files_list = [
        "index.html — головна сторінка з навігацією, метою та постановкою завдань;",
        "task1.html — інтерфейс Завдання 1 (Кіноафіша);",
        "task2.html — інтерфейс Завдання 2 (DOM-маніпуляції та числова таблиця);",
        "CSS/style.css — каскадна таблиця стилів інтерфейсу, таблиць та кнопок;",
        "JS/task1.js — сценарій опрацювання кіноафіші та дат;",
        "JS/task2.js — сценарій обходу DOM-вузлів та математичної обробки таблиці."
    ]
    for f in files_list:
        p_f = doc.add_paragraph(style='List Bullet')
        p_f.paragraph_format.line_spacing = 1.25
        p_f.paragraph_format.space_after = Pt(2)
        r = p_f.add_run(f)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)

    # Функція вставки ілюстрацій
    img_folder = r"C:\Users\SIN\WebGit\LR5\report_screenshots"
    
    def add_figure(filename, caption_text, width=Inches(6.2)):
        img_path = os.path.join(img_folder, filename)
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.add_run().add_picture(img_path, width=width)
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(12)
            r = p_cap.add_run(caption_text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)

    add_figure("fig01_index.png", "Рис. 1 - Головна сторінка веб-застосунку з навігаційним меню та описом завдань")
    add_figure("fig02_task1_initial.png", "Рис. 2 - Початковий вигляд інтерфейсу Завдання 1 (Кіноафіша) перед запуском скрипта")

    doc.add_page_break()

    # Секція 1: Кіноафіша
    add_p("1. Реалізація кіноафіші та обробка дат у JavaScript (Завдання 1)", bold=True, size=15, space_before=6, space_after=6)
    add_p("Вихідні дані про сеанси організовано у вигляді масиву об'єктів arrMovies, який містить атрибути Name, Cinema, Hall, Data, Time. Для всебічної перевірки алгоритму частина фільмів має порожнє значення поля Hall.")
    add_p("Для розрахунку різниці між поточною датою та датою прем'єри реалізовано функцію ras_Diff_Days(text). Вона здійснює парсинг рядка формату DD.MM.YYYY, конструює об'єкт Date, обчислює різницю у мілісекундах відносно поточної дати (new Date()) та конвертує її у кількість повних днів: Math.floor(timeDiff / (1000 * 3600 * 24)).")
    add_p("Функція vivod() здійснює обхід елементів масиву та класифікує статус прокату кожного фільму: очікується прем'єра (білий фон), активний прокат (зелений фон #c8e6c9), завершення прокату за <= 2 дні (попереджувальний жовтий фон #fff9c4) та завершений показ (червоний фон #ffcdd2). Якщо зал не вказано, динамічно вставляється примітка «Не вказано (у кінотеатрі лише один зал)», як того вимагає варіант 6.")

    add_figure("fig03_task1_table.png", "Рис. 3 - Динамічно згенерована кіноафіша з підрахунком днів, кольоровими статусами та перевіркою залів", width=Inches(6.2))

    doc.add_page_break()

    # Секція 2: Завдання 2.1
    add_p("2. Обхід вузлів DOM та формування нумерованого списку (Завдання 2.1)", bold=True, size=15, space_before=6, space_after=6)
    add_p("У Завданні 2.1 продемонстровано взаємодію з об'єктною моделлю документа (Document Object Model). На веб-сторінці розташовано низку текстових блоків <div class=\"source-block\">, присвячених структурі DOM-дерева.")
    add_p("Функція div_to_list() отримує колекцію зазначених елементів за допомогою методу document.getElementsByClassName(\"source-block\"). Потім у циклі for здійснюється зчитування чистого тексту з кожного вузла через властивість textContent та генерація HTML-рядка нумерованого списку <ol>, який монтується у документ за допомогою innerHTML.")

    add_figure("fig04_task2_initial.png", "Рис. 4 - Загальний вигляд сторінки Завдання 2 з вихідними текстовими блоками та матрицею чисел", width=Inches(6.0))
    add_figure("fig05_task2_1_list.png", "Рис. 5 - Сформований нумерований список після обходу блоків div засобами DOM (Завдання 2.1)", width=Inches(6.0))

    doc.add_page_break()

    # Секція 3: Завдання 2.2
    add_p("3. Обробка числової матриці, пошук екстремумів та стилізація (Завдання 2.2)", bold=True, size=15, space_before=6, space_after=6)
    add_p("У Завданні 2.2 реалізовано зчитування числової таблиці розміром 4х4 у двовимірний масив засобами функції generate_arr(). Зчитування відбувається шляхом доступу до колекцій table.rows[i].cells[j] та парсингу рядків через parseFloat().")
    add_p("Функція table_obrab() виконує пошук екстремальних значень окремо для додатних (maxPos, minPos) та від'ємних елементів (maxNeg — найближче до 0, minNeg — найменше від'ємне число). На другому етапі відбувається динамічне розфарбовування чарунок через cell.style.backgroundColor: додатні позначаються блакитним (#bbdefb), максимальне додатне — яскраво-зеленим (#2ecc71), мінімальне додатне — м'ятним (#a8e6cf); від'ємні — помаранчевим (#ffe0b2), мінімальне від'ємне — контрастним червоним (#e74c3c), максимальне від'ємне — кораловим (#ffab91).")

    add_figure("fig04b_task2_2_initial.png", "Рис. 6 - Вихідна числова таблиця 4х4 до виконання алгоритму обробки (Завдання 2.2)", width=Inches(5.8))
    add_figure("fig06b_task2_2_colored.png", "Рис. 7 - Результат пошуку екстремумів та кольорового маркування чарунок матриці (Завдання 2.2, Варіант 6)", width=Inches(5.8))

    doc.add_page_break()

    # Секція 4: Додаткове завдання та Висновок
    add_p("4. Додаткове завдання: розрахунок сум елементів по стовпцях", bold=True, size=15, space_before=6, space_after=6)
    add_p("У функції col_sums() реалізовано математичний алгоритм підсумовування елементів матриці за кожним окремим вертикальним стовпцем, що демонструє роботу із вкладеними циклами та індексацією матриць.")

    add_figure("fig07_task2_2_cols.png", "Рис. 8 - Результат підрахунку сум числових елементів за стовпцями матриці (Завдання 2.2)", width=Inches(5.8))

    add_p("Висновок", bold=True, size=16, space_before=16, space_after=8)
    add_p("У ході виконання лабораторної роботи №5 було успішно освоєно базові та поглиблені принципи програмування на стороні клієнта мовою JavaScript, роботу з об'єктною моделлю документа (DOM) та обробку подій користувача.")
    add_p("Під час виконання Завдання 1 опрацьовано механізми маніпуляції з об'єктом Date(): виконано коректне перетворення текстових форматів дат на об'єкти часу, розрахунок інтервалів у мілісекундах та їх конвертацію в календарні дні. Засвоєно структурування комплексних сутностей через масиви об'єктів, методи ітерації (forEach), перевірку граничних і порожніх значень та динамічну генерацію табличних структур.")
    add_p("У Завданні 2.1 освоєно техніку обходу вузлів HTML-документа за допомогою інтерфейсу DOM (getElementsByClassName), безпечне вилучення текстового вмісту через властивість textContent та конструювання впорядкованих списків.")
    add_p("У Завданні 2.2 закріплено алгоритми обробки двовимірних матриць у JavaScript: вилучення числових даних із чарунок таблиці, пошук екстремумів серед множин додатних і від'ємних чисел, а також інтерактивне керування візуальним оформленням сторінки через модифікацію інлайн-стилів елементів (style.backgroundColor, style.color). Додатково розроблено алгоритм вертикального підсумовування елементів по стовпцях.")
    add_p("Усі розроблені модулі протестовано у сучасному веб-браузері. Код структуровано за модульними стандартами із розділенням структури, стилів та логіки, що забезпечує високу якість, підтримуваність та надійність створеного рішення.")

    out_docx = r"C:\Users\SIN\WebGit\LR5\ПРВЗ_ЛР№5_Вар№16_Мамон.docx"
    doc.save(out_docx)
    print(f"Saved DOCX to {out_docx}")

if __name__ == '__main__':
    create_report()
